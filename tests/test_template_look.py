"""Drafts copy the tagged template's look. Templates here are made up in the test — never a real resume."""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

import resume  # noqa: E402
import template_look as tl  # noqa: E402

TMP = tempfile.mkdtemp()
PROFILE = {"full_name": "Pat Example", "email": "pat@example.com", "phone": "555-010-0199", "city_state": "Springfield, ST"}


def _border(p, color="AA0000"):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "double"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def _tab(p, pos=10800):
    tabs = OxmlElement("w:tabs")
    t = OxmlElement("w:tab")
    t.set(qn("w:val"), "right")
    t.set(qn("w:pos"), str(pos))
    tabs.append(t)
    p._p.get_or_add_pPr().append(tabs)


def _wingdings_bullets(doc):
    """Point every bullet level at a Wingdings arrow, like many real resumes do."""
    num = doc.part.numbering_part.element
    for lvl in num.xpath("./w:abstractNum/w:lvl[w:numFmt/@w:val='bullet']"):
        lvl.find(qn("w:lvlText")).set(qn("w:val"), "")
        rpr = lvl.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            lvl.append(rpr)
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi"):
            rf.set(qn(a), "Wingdings")


def make_template(path, table=False, two_cols=False, headings=("Career Summary", "Areas of Expertise",
                                                                  "Work History", "Education")):
    d = Document()
    p = d.add_paragraph()
    _tab(p)
    r = p.add_run("Sample Person")
    r.bold, r.font.size = True, Pt(20)
    r = p.add_run("\tsample@example.org | 555-000-1111")
    r.font.size = Pt(9)

    def heading(text):
        h = d.add_paragraph()
        _border(h)
        hr = h.add_run(text.upper())
        hr.font.color.rgb = RGBColor.from_string("AA0000")
        hr.font.size = Pt(13)
        hr.font.small_caps = True
        return h
    heading(headings[0])
    d.add_paragraph("Operations leader with broad experience running large teams across many places.")
    d.add_paragraph("Led a big program", style="List Bullet")
    heading(headings[1])
    d.add_paragraph("Planning | Budgets | Teams")
    heading(headings[2])
    o = d.add_paragraph()
    _tab(o)
    o.add_run("Example Corp").bold = True
    dr = o.add_run("\tJan 2020 - Present")
    dr.italic = True
    t = d.add_paragraph()
    tr = t.add_run("Director of Things")
    tr.font.color.rgb = RGBColor.from_string("006600")
    tr.underline = True
    s = d.add_paragraph()
    s.add_run("Ran a team of people across several sites with a large yearly budget for the company.").italic = True
    fb = d.add_paragraph(style="List Bullet")
    lr = fb.add_run("Key Win: ")
    lr.bold = True
    fb.add_run("Delivered the main thing.")
    d.add_paragraph("Did another thing.", style="List Bullet")
    heading(headings[3])
    d.add_paragraph("BS, Example University")
    if table:
        d.add_table(rows=1, cols=2)
    if two_cols:
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), "2")
        d.sections[0]._sectPr.append(cols)
    _wingdings_bullets(d)
    d.save(path)
    return path


CONTENT = resume.normalize_resume({
    "headline": "Operations Executive", "summary": "Leads large teams.", "highlights": ["Did a big thing"],
    "credentials_line": "", "competencies": ["Planning", "Budgets", "Risk"],
    "experience": [{"org": "New Org", "dates": "Jan 2021 - Present",
                    "roles": [{"title": "Chief Operating Officer", "dates": "", "summary": "Runs operations.",
                               "flagship": "Cut costs.", "bullets": ["Built a team", "Opened a site"]}]}],
    "additional_roles": [{"title": "Manager, Old Org", "dates": "Jan 2010 - Dec 2012", "bullets": []}],
    "education": ["MBA, Some School"], "certifications": "PMP", "recognitions": "", "technology_skills": ""})


def para(doc, text):
    return next(p for p in doc.paragraphs if p.text.strip() == text or p.text.startswith(text))


def rpr_xml(run):
    el = run._r.find(qn("w:rPr"))
    return "" if el is None else el.xml


class Learn(unittest.TestCase):
    def test_parts_found(self):
        look, why = tl.learn(Document(make_template(os.path.join(TMP, "t1.docx"))))
        self.assertEqual(why, "")
        for part in ("name", "heading", "summary", "highlight", "org", "title", "scope", "flagship", "bullet"):
            self.assertIn(part, look.parts, part)
        self.assertTrue(look.name_with_contact)
        self.assertEqual([k for k, _ in look.sections], ["summary", "competencies", "experience", "education"])
        self.assertEqual(look.parts["flagship"]["label"][0], "Key Win: ")

    def test_unsupported_templates_fall_back(self):
        for kw, words in (({"table": True}, "table"), ({"two_cols": True}, "column")):
            path = make_template(os.path.join(TMP, f"t-{words}.docx"), **kw)
            out = os.path.join(TMP, f"o-{words}.docx")
            info = resume.build_docx(CONTENT, PROFILE, path, out)
            self.assertEqual(info["mode"], "built-in")
            self.assertIn(words, info["reason"])
            self.assertTrue(os.path.exists(out))

    def test_no_experience_heading_falls_back(self):
        path = make_template(os.path.join(TMP, "t-noexp.docx"), headings=("Summary", "Skills", "Stuff", "Education"))
        info = resume.build_docx(CONTENT, PROFILE, path, os.path.join(TMP, "o-noexp.docx"))
        self.assertEqual(info["mode"], "built-in")
        self.assertIn("Experience", info["reason"])


class Build(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tpl = make_template(os.path.join(TMP, "t2.docx"))
        cls.out = os.path.join(TMP, "o2.docx")
        cls.info = resume.build_docx(CONTENT, PROFILE, cls.tpl, cls.out)
        cls.t = Document(cls.tpl)
        cls.d = Document(cls.out)

    def test_mode_and_section_names(self):
        self.assertEqual(self.info["mode"], "template")
        texts = [p.text for p in self.d.paragraphs]
        self.assertEqual([x for x in texts if x.isupper() and len(x) > 3],
                         ["CAREER SUMMARY", "AREAS OF EXPERTISE", "WORK HISTORY", "EDUCATION"])
        self.assertNotIn("Professional Experience", " ".join(texts))
        self.assertFalse(any("Sample Person" in x or "Example Corp" in x for x in texts))   # no template words left

    def test_heading_look_copied(self):
        th, dh = para(self.t, "WORK HISTORY"), para(self.d, "WORK HISTORY")
        self.assertEqual(rpr_xml(dh.runs[0]), rpr_xml(th.runs[0]))
        self.assertIn('w:val="double"', dh._p.pPr.xml)

    def test_job_lines_copied(self):
        org = para(self.d, "New Org")
        self.assertTrue(org.runs[0].bold)
        self.assertTrue(org.runs[1].italic)          # dates look like the template's dates
        self.assertIn("w:tabs", org._p.pPr.xml)
        title = para(self.d, "Chief Operating Officer")
        self.assertEqual(rpr_xml(title.runs[0]), rpr_xml(para(self.t, "Director of Things").runs[0]))
        self.assertTrue(para(self.d, "Runs operations.").runs[0].italic)

    def test_bullets_and_flagship(self):
        b = para(self.d, "Built a team")
        self.assertEqual(b.style.name, "List Bullet")   # bullet style (and its numbering) carried over
        f = para(self.d, "Key Win: Cut costs.")
        self.assertTrue(f.runs[0].bold)
        self.assertFalse(bool(f.runs[1].bold))
        num = self.d.part.numbering_part.element.xml
        self.assertNotIn("Wingdings", num)
        self.assertNotIn("", num)

    def test_name_line_and_checks(self):
        first = self.d.paragraphs[0]
        self.assertTrue(first.text.startswith("Pat Example\t"))
        self.assertIn("pat@example.com", first.text)
        self.assertEqual(first.runs[0].font.size, Pt(20))
        ats = resume.ats_check(self.out, PROFILE)
        fails = [c["name"] for c in ats["checks"] if c["level"] == "fail"]
        self.assertFalse([f for f in fails if "heading" in f or "font" in f or "Name" in f], fails)

    def test_plain_text_uses_template_names(self):
        txt = resume.to_plain_text(CONTENT, PROFILE, self.info["names"])
        self.assertIn("WORK HISTORY", txt)
        self.assertIn("AREAS OF EXPERTISE", txt)

    def test_no_template_uses_builtin(self):
        info = resume.build_docx(CONTENT, PROFILE, "", os.path.join(TMP, "o3.docx"))
        self.assertEqual(info["mode"], "built-in")


class Detect(unittest.TestCase):
    def test_section_words(self):
        for text, kind in (("PROFESSIONAL EXPERIENCE", "experience"), ("Work History", "experience"),
                           ("Core Competencies", "competencies"), ("Education & Certifications", "education"),
                           ("Technology Skills", "tech"), ("Awards", "recognitions"),
                           ("Additional Executive Leadership:", "additional"), ("Chief Operating Officer", None),
                           ("Led a large program across many sites for the company", None)):
            self.assertEqual(tl.section_kind(text), kind, text)


if __name__ == "__main__":
    unittest.main()
