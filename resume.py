"""Resume drafting engine: prompt, fact check, ATS-safe .docx builder, ATS self-check, JD fetch."""
import copy
import html
import re
import urllib.parse
import urllib.request

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ACCENT = "1E3764"          # navy used in the source resume
HEADING_FILL = "F2F2F2"    # light gray section band
BULLET_CHAR = "•"     # standard bullet; ATS-safe (no Wingdings/Symbol)
BODY_FONT = "Calibri"
BODY_SIZE = 10.5

# ---------------------------------------------------------------- text helpers

def extract_docx_text(path):
    """Plain text of a .docx (paragraphs, in order). Used as source material + ATS check."""
    d = Document(path)
    lines = []
    for p in d.paragraphs:
        t = p.text.replace("\t", "  ").strip()
        if t:
            lines.append(t)
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def fetch_posting(url, timeout=20):
    """Best-effort fetch of a job posting's text. Returns (ok, text, reason)."""
    if not url or not url.startswith("http"):
        return False, "", "No link on this job."
    host = (urllib.parse.urlparse(url).hostname or "").lower()
    if host == "indeed.com" or host.endswith(".indeed.com"):
        return False, "", "Indeed blocks automated reading. Open the posting, copy the description, and paste it here."
    req = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.0 Safari/605.1.15",
        "Accept": "text/html,application/xhtml+xml",
    })
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            raw = r.read(3_000_000).decode(r.headers.get_content_charset() or "utf-8", "replace")
    except Exception as e:  # noqa: BLE001 — any failure means "paste it"
        return False, "", f"Couldn't open the posting ({e.__class__.__name__}). Paste the job description instead."
    # JSON-LD JobPosting (Workday, Greenhouse, many ATS) — best source when present
    m = re.search(r'"description"\s*:\s*"(.{400,}?)"\s*,\s*"', raw, re.S)
    text = ""
    if m and "JobPosting" in raw:
        text = m.group(1).encode("utf-8").decode("unicode_escape", "ignore")
    if len(text) < 400:
        body = re.sub(r"(?is)<(script|style|noscript|svg|head)[^>]*>.*?</\1>", " ", raw)
        body = re.sub(r"(?i)<br\s*/?>|</(p|li|div|h\d)>", "\n", body)
        text = re.sub(r"<[^>]+>", " ", body)
    text = html.unescape(re.sub(r"<[^>]+>", " ", text))
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text).strip()
    if len(text) < 800:
        return False, text, "The posting page loads its text with JavaScript, so it couldn't be read automatically. Paste the job description instead."
    return True, text[:15000], ""


# ---------------------------------------------------------------- prompt

RESUME_SCHEMA = """{
  "headline": "one line, e.g. 'Senior Technology & Cybersecurity Executive | AI Governance | Cloud & Enterprise Risk'",
  "summary": "3-4 sentence professional summary tailored to this job",
  "highlights": ["3-4 top accomplishment bullets aimed at this job"],
  "credentials_line": "short bold line, e.g. clearance | notable distinction (or empty)",
  "competencies": ["12-18 short competency phrases, using the job's own vocabulary where truthful"],
  "experience": [
    {"org": "Employer or service line, e.g. 'Acme Systems Inc.' or 'U.S. Army'",
     "dates": "Mon YYYY - Mon YYYY or Present",
     "roles": [
       {"title": "Role title with civilian equivalent in parentheses when helpful",
        "dates": "only for multi-role orgs, else empty",
        "summary": "1-2 sentence scope statement (people, budget, users, sites)",
        "flagship": "single strongest accomplishment for THIS job (or empty)",
        "bullets": ["2-4 accomplishment bullets"]}
     ]}
  ],
  "additional_roles": [{"title": "Title, Organization", "dates": "Mon YYYY - Mon YYYY", "bullets": ["0-1 bullet"]}],
  "education": ["Degree, Institution"],
  "certifications": "PMP | CISSP | ...",
  "recognitions": "one line (or empty)",
  "technology_skills": "one line of pipe-separated skills (or empty)",
  "notes_for_candidate": ["what you emphasized and why", "gaps vs. the posting's requirements", "anything to verify"]
}"""

RESUME_SYSTEM = """You are an expert executive resume writer who specializes in translating senior military and government leadership into civilian corporate language.

Write a tailored two-page resume for ONE specific job, using ONLY facts from the candidate's source material.

Hard rules:
1. TRUTH: never invent employers, titles, dates, degrees, certifications, clearances, numbers, percentages, dollar amounts, headcounts or outcomes. Every number you write must appear in the source material. If the job wants something the candidate lacks, do NOT claim it — list it under notes_for_candidate as a gap.
2. TRANSLATE: express military roles in civilian terms (e.g. "Deputy Commanding Officer (Chief Operating Officer)"), replace military jargon and acronyms with plain business language unless the employer is defense/government and the acronym is standard there.
3. TAILOR: select, reorder and rephrase the candidate's real accomplishments to match the job's priorities and vocabulary. Lead with what this employer cares about most.
4. ATS: plain text only — no emojis, no special symbols, no tables. Standard section names. Dates as "Mon YYYY - Mon YYYY".
5. LENGTH: two pages, roughly 650-850 words total. Most recent and most relevant roles get the most detail; older roles go in additional_roles with 0-1 bullet.
6. Keep the candidate's employment history complete and in reverse-chronological order.

Return ONLY a JSON object matching this schema (no commentary, no markdown fences):
""" + RESUME_SCHEMA


def build_resume_prompt(profile, inventory, resume_texts, job, jd_text, template_name=""):
    parts = ["# CANDIDATE PROFILE",
             f"Name: {profile.get('full_name')}",
             f"Location: {profile.get('city_state')}",
             f"Clearance: {profile.get('clearance') or 'not specified'}"]
    prefs = [(lbl, profile.get(k, "")) for k, lbl in (("levels", "Target levels"), ("lanes", "Target role lanes"))]
    prefs = [f"{lbl}: {v}" for lbl, v in prefs if str(v).strip()]
    if prefs:
        parts += prefs
    if str(profile.get("notes", "")).strip():
        parts += ["", "# CANDIDATE'S INSTRUCTIONS TO THE WRITER (follow these; they override style defaults, never the truth rules)",
                  profile["notes"].strip()]
    if inventory.strip():
        parts += ["", "# CAREER INVENTORY (primary source of facts)", inventory.strip()]
    # The resume tagged as the layout template is the current one: follow it. Older ones are for facts only.
    current = [(n, t) for n, t in resume_texts if template_name and n == template_name]
    older = [(n, t) for n, t in resume_texts if not (template_name and n == template_name)]
    for name, text in current:
        parts += ["", f"# CURRENT RESUME — {name} (follow its section names, section order, structure and wording style)",
                  text.strip()]
    for name, text in older:
        parts += ["", f"# {'OLDER' if current else 'EXISTING'} RESUME — {name}"
                      + (" (use only for facts; do not copy its structure or wording)" if current else ""), text.strip()]
    parts += ["", "# TARGET JOB",
              f"Title: {job.get('title')}",
              f"Company: {job.get('company')}",
              f"Location: {job.get('location')} ({job.get('work_mode')})",
              f"Clearance noted: {job.get('clearance') or 'n/a'}",
              f"Role lane: {job.get('lane') or 'n/a'}",
              f"Why it was flagged as a fit: {job.get('fit_reason') or 'n/a'}",
              "", "# JOB DESCRIPTION",
              (jd_text or "(not available — tailor to the title and company above)").strip()[:15000],
              "", "Write the tailored resume JSON now."]
    return RESUME_SYSTEM, "\n".join(parts)


def normalize_resume(c):
    """Fill defaults / coerce types so the builder never crashes on sloppy model output."""
    def s(v):
        return "" if v is None else str(v).strip()

    def ls(v):
        if isinstance(v, str):
            v = [x for x in re.split(r"\n|;\s", v) if x.strip()]
        return [s(x) for x in (v or []) if s(x)]
    out = {
        "headline": s(c.get("headline")),
        "summary": s(c.get("summary")),
        "highlights": ls(c.get("highlights"))[:5],
        "credentials_line": s(c.get("credentials_line")),
        "competencies": ls(c.get("competencies"))[:20],
        "experience": [],
        "additional_roles": [],
        "education": ls(c.get("education")),
        "certifications": s(c.get("certifications")) if not isinstance(c.get("certifications"), list)
        else " | ".join(ls(c.get("certifications"))),
        "recognitions": s(c.get("recognitions")) if not isinstance(c.get("recognitions"), list)
        else "; ".join(ls(c.get("recognitions"))),
        "technology_skills": s(c.get("technology_skills")) if not isinstance(c.get("technology_skills"), list)
        else " | ".join(ls(c.get("technology_skills"))),
        "notes_for_candidate": ls(c.get("notes_for_candidate")),
    }
    for e in c.get("experience") or []:
        if not isinstance(e, dict):
            continue
        roles = []
        for r in e.get("roles") or []:
            if isinstance(r, dict):
                roles.append({"title": s(r.get("title")), "dates": s(r.get("dates")), "summary": s(r.get("summary")),
                              "flagship": s(r.get("flagship")), "bullets": ls(r.get("bullets"))[:6]})
        out["experience"].append({"org": s(e.get("org")), "dates": s(e.get("dates")), "roles": roles})
    for r in c.get("additional_roles") or []:
        if isinstance(r, dict):
            out["additional_roles"].append({"title": s(r.get("title")), "dates": s(r.get("dates")),
                                            "bullets": ls(r.get("bullets"))[:2]})
    return out


# ---------------------------------------------------------------- fact check

_NUM_RE = re.compile(r"\$?\d[\d,]*(?:\.\d+)?\s?(?:%|[KMB]\b|\+|x\b)?", re.I)


def _norm_num(tok):
    t = tok.lower().replace(",", "").replace("$", "").replace(" ", "").rstrip("+")
    return t


_CODE_RE = re.compile(r"\b(?:[A-Z]{2,}[ -]?)?\d+(?:-\d+)+\b|\b(?:SP|ISO|NIST|IL|FIPS|DoDI?|AR|CMMC Level|Level|Tier|Title|GS|SES)[ -]?\d+\b")


def fact_check(content, source_text):
    """Numbers in the draft that don't appear anywhere in the source material.

    Standard identifiers (NIST SP 800-37, ISO 27001, GS-15, IL5 ...) are ignored — they aren't claims.
    """
    src = {_norm_num(m.group()) for m in _NUM_RE.finditer(source_text)}
    src_digits = {re.sub(r"\D", "", x) for x in src}
    flagged = []
    for line in content_lines(content):
        for m in _NUM_RE.finditer(_CODE_RE.sub(" ", line)):
            tok = m.group().strip()
            n = _norm_num(tok)
            digits = re.sub(r"\D", "", n)
            if not digits or (len(digits) == 4 and digits.startswith(("19", "20"))):
                continue  # years
            if n in src or digits in src_digits:
                continue
            flagged.append({"value": tok, "context": line[:160]})
    seen, uniq = set(), []
    for f in flagged:
        if f["value"] not in seen:
            seen.add(f["value"])
            uniq.append(f)
    return uniq


def content_lines(c):
    lines = [c["headline"], c["summary"], *c["highlights"], c["credentials_line"], " | ".join(c["competencies"])]
    for e in c["experience"]:
        lines.append(f"{e['org']} {e['dates']}")
        for r in e["roles"]:
            lines += [f"{r['title']} {r['dates']}", r["summary"], r["flagship"], *r["bullets"]]
    for r in c["additional_roles"]:
        lines += [f"{r['title']} {r['dates']}", *r["bullets"]]
    lines += [*c["education"], c["certifications"], c["recognitions"], c["technology_skills"]]
    return [x for x in lines if x]


def contact_parts(profile):
    return [x for x in (profile.get("city_state"), profile.get("phone"), profile.get("email"),
                        profile.get("linkedin")) if x]


def to_plain_text(c, profile):
    L = [profile.get("full_name", ""), " | ".join(contact_parts(profile)), ""]
    if c["headline"]:
        L.append(c["headline"])
    if c["summary"]:
        L.append(c["summary"])
    L += [f"- {h}" for h in c["highlights"]]
    L += ["", "CORE COMPETENCIES AND SKILLS"]
    if c["credentials_line"]:
        L.append(c["credentials_line"])
    if c["competencies"]:
        L.append(" | ".join(c["competencies"]))
    L += ["", "PROFESSIONAL EXPERIENCE"]
    for e in c["experience"]:
        L.append(f"{e['org']}    {e['dates']}".rstrip())
        for r in e["roles"]:
            L.append(r["title"] + (f" | {r['dates']}" if r["dates"] else ""))
            if r["summary"]:
                L.append(r["summary"])
            if r["flagship"]:
                L.append(f"- Flagship Accomplishment: {r['flagship']}")
            L += [f"- {b}" for b in r["bullets"]]
        L.append("")
    if c["additional_roles"]:
        L.append("Additional Executive Leadership:")
        for r in c["additional_roles"]:
            L.append(r["title"] + (f" | {r['dates']}" if r["dates"] else ""))
            L += [f"- {b}" for b in r["bullets"]]
        L.append("")
    L.append("EDUCATION & CERTIFICATIONS")
    L += c["education"]
    if c["certifications"]:
        L.append(c["certifications"])
    if c["recognitions"]:
        L += ["", "RECOGNITIONS", c["recognitions"]]
    if c["technology_skills"]:
        L += ["", "TECHNOLOGY SKILLS", c["technology_skills"]]
    return "\n".join(L).strip() + "\n"


# ---------------------------------------------------------------- docx builder

def _twips(pt):
    return str(int(round(pt * 20)))


def _pPr(p):
    return p._p.get_or_add_pPr()


def _set_spacing(p, before=0, after=0, line=None):
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), _twips(before))
    sp.set(qn("w:after"), _twips(after))
    if line:
        sp.set(qn("w:line"), str(line))
        sp.set(qn("w:lineRule"), "auto")
    _pPr(p).append(sp)


def _right_tab(p, pos_twips):
    tabs = OxmlElement("w:tabs")
    t = OxmlElement("w:tab")
    t.set(qn("w:val"), "right")
    t.set(qn("w:pos"), str(pos_twips))
    tabs.append(t)
    _pPr(p).append(tabs)


def _bottom_border(p, sz=12, color=ACCENT, space=1):
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), color)
    bdr.append(b)
    _pPr(p).append(bdr)


def _shade(p, fill=HEADING_FILL):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _pPr(p).append(shd)


def _run(p, text, bold=False, italic=False, size=None, color=None, underline=False):
    r = p.add_run(text)
    r.bold = bold or None
    r.italic = italic or None
    if underline:
        r.underline = True
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _ensure_bullet_numbering(doc):
    """Create two ATS-safe bullet list definitions (navy for highlights, black for body). Returns (hi_numId, body_numId)."""
    numbering = doc.part.numbering_part.element
    ids = [int(x) for x in numbering.xpath("./w:abstractNum/@w:abstractNumId")] or [0]
    nids = [int(x) for x in numbering.xpath("./w:num/@w:numId")] or [0]
    result = []
    for i, (color, left, hanging) in enumerate(((ACCENT, 630, 270), (None, 720, 360))):
        aid = max(ids) + 1 + i
        nid = max(nids) + 1 + i
        an = OxmlElement("w:abstractNum")
        an.set(qn("w:abstractNumId"), str(aid))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "singleLevel")
        an.append(mlt)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", BULLET_CHAR), ("w:lvlJc", "left")):
            el = OxmlElement(tag)
            el.set(qn("w:val"), val)
            lvl.append(el)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        ppr.append(ind)
        lvl.append(ppr)
        rpr = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(a), BODY_FONT)
        rpr.append(rf)
        if color:
            c = OxmlElement("w:color")
            c.set(qn("w:val"), color)
            rpr.append(c)
        lvl.append(rpr)
        an.append(lvl)
        # abstractNum elements must precede num elements
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(an)
        else:
            numbering.append(an)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(nid))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(aid))
        num.append(ref)
        numbering.append(num)
        result.append(nid)
    return tuple(result)


def _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, zero_indent=True):
    p = doc.add_paragraph()
    p.alignment = align
    if zero_indent:  # template "Normal" styles often carry an indent; the source resume overrides it per paragraph
        ind = OxmlElement("w:ind")
        for a in ("w:left", "w:right", "w:firstLine"):
            ind.set(qn(a), "0")
        _pPr(p).append(ind)
    return p


def _bullet(doc, num_id, text, after=0, prefix=None, italic=False, shaded=False):
    p = _para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, zero_indent=False)
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), "0")
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), str(num_id))
    numPr.append(il)
    numPr.append(ni)
    _pPr(p).append(numPr)
    _set_spacing(p, 0, after)
    if shaded:
        _shade(p)
    if prefix:
        _run(p, prefix, italic=italic)
    _run(p, text, italic=italic)
    return p


def build_docx(c, profile, template_path, out_path):
    """Build the resume .docx. Uses the candidate's own .docx for page setup, fonts, footers and styles."""
    if template_path:
        doc = Document(template_path)
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
    else:
        doc = Document()
        s = doc.sections[0]
        from docx.shared import Inches
        s.left_margin = s.right_margin = Inches(0.5)
        s.top_margin = s.bottom_margin = Inches(0.5)
    normal = doc.styles["Normal"]
    normal.font.name = normal.font.name or BODY_FONT
    if not normal.font.size:
        normal.font.size = Pt(BODY_SIZE)
    sec = doc.sections[0]
    usable = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)  # EMU -> twips
    hi_num, body_num = _ensure_bullet_numbering(doc)

    # Name + contact (in the body, not a header — many ATS skip headers)
    p = _para(doc)
    _right_tab(p, usable)
    _set_spacing(p, 0, 2)
    _run(p, profile.get("full_name") or "Your Name", size=16)
    _run(p, "\t" + " | ".join(contact_parts(profile)), size=11)
    _bottom_border(p, sz=12, space=4)

    # Headline + summary
    p = _para(doc)
    _set_spacing(p, 6, 6)
    if c["headline"]:
        _run(p, c["headline"], bold=True)
        p.add_run().add_break()
    _run(p, c["summary"])
    for h in c["highlights"]:
        _bullet(doc, hi_num, h)

    def heading(text):
        hp = _para(doc)
        _set_spacing(hp, 8, 5.3)
        _bottom_border(hp, sz=12, space=0)
        _shade(hp)
        _run(hp, text)

    heading("Core Competencies and Skills")
    if c["credentials_line"]:
        p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER)
        _set_spacing(p, 0, 2)
        _run(p, c["credentials_line"], bold=True)
    if c["competencies"]:
        p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER)
        _set_spacing(p, 0, 2)
        _run(p, " | ".join(c["competencies"]))

    heading("Professional Experience")
    for i, e in enumerate(c["experience"]):
        p = _para(doc)
        _right_tab(p, usable)
        _set_spacing(p, 6 if i else 0, 1.1)
        _run(p, e["org"], bold=True)
        if e["dates"]:
            _run(p, "\t" + e["dates"])
        for j, r in enumerate(e["roles"]):
            if r["title"]:
                p = _para(doc)
                _set_spacing(p, 4 if j else 0, 0)
                _run(p, r["title"], bold=True, color=ACCENT)
                if r["dates"]:
                    _run(p, f" | {r['dates']}", color=ACCENT)
            if r["summary"]:
                p = _para(doc)
                _set_spacing(p, 0, 1.1)
                _run(p, r["summary"], italic=True)
            if r["flagship"]:
                _bullet(doc, hi_num, r["flagship"], prefix="Flagship Accomplishment: ", italic=True, shaded=True)
            for b in r["bullets"]:
                _bullet(doc, body_num, b)

    if c["additional_roles"]:
        p = _para(doc)
        _set_spacing(p, 6, 1)
        _run(p, "Additional Executive Leadership:", bold=True)
        for r in c["additional_roles"]:
            p = _para(doc)
            _set_spacing(p, 0, 0)
            _run(p, r["title"] + (f" | {r['dates']}" if r["dates"] else ""), bold=True)
            for b in r["bullets"]:
                _bullet(doc, body_num, b)

    heading("Education & Certifications")
    for ed in c["education"]:
        p = _para(doc)
        _set_spacing(p, 0, 0)
        _run(p, ed)
    if c["certifications"]:
        p = _para(doc)
        _set_spacing(p, 0, 0)
        _run(p, c["certifications"])
    if c["recognitions"]:
        heading("Recognitions")
        p = _para(doc)
        _set_spacing(p, 0, 0)
        _run(p, c["recognitions"])
    if c["technology_skills"]:
        heading("Technology Skills")
        p = _para(doc)
        _set_spacing(p, 0, 0)
        _run(p, c["technology_skills"])
    doc.core_properties.title = f"{profile.get('full_name', '')} - Resume".strip(" -")
    doc.core_properties.author = profile.get("full_name", "")
    doc.save(out_path)


# ---------------------------------------------------------------- ATS self-check

_BAD_FONTS = ("wingdings", "symbol", "webdings", "zapf")


def ats_check(docx_path, profile):
    d = Document(docx_path)
    body = d.element.body
    xml = body.xml if hasattr(body, "xml") else ""
    text = "\n".join(p.text for p in d.paragraphs)
    checks = []

    def add(name, ok, detail="", warn=False):
        checks.append({"name": name, "ok": bool(ok), "level": "warn" if (warn and not ok) else ("ok" if ok else "fail"),
                       "detail": detail})

    first = d.paragraphs[0].text if d.paragraphs else ""
    name = profile.get("full_name", "")
    add("Name at top of page (in body, not header)", bool(name) and name in first, first[:80])
    for label, key in (("Email", "email"), ("Phone", "phone")):
        v = profile.get(key, "")
        add(f"{label} readable in body text", bool(v) and v in text, v or "not set in Profile")
    for h in ("Professional Experience", "Education"):
        add(f"Standard section heading: {h}", h.lower() in text.lower())
    n_tables = len(body.findall(qn("w:tbl")))
    n_txbx = len(body.xpath(".//w:txbxContent"))
    n_draw = len(body.xpath(".//w:drawing")) + len(body.xpath(".//w:pict"))
    add("No tables", n_tables == 0, f"{n_tables} found")
    add("No text boxes", n_txbx == 0, f"{n_txbx} found")
    add("No images or shapes", n_draw == 0, f"{n_draw} found")
    fonts = {f.lower() for f in body.xpath(".//w:rFonts/@w:ascii")}
    num_xml = d.part.numbering_part.element if d.part._rels else None
    used_num = set(body.xpath(".//w:numPr/w:numId/@w:val"))
    bullet_fonts = set()
    if num_xml is not None:
        for nid in used_num:
            aid = num_xml.xpath(f"./w:num[@w:numId='{nid}']/w:abstractNumId/@w:val")
            if aid:
                bullet_fonts |= {f.lower() for f in num_xml.xpath(
                    f"./w:abstractNum[@w:abstractNumId='{aid[0]}']//w:rFonts/@w:ascii")}
    bad = sorted(f for f in fonts | bullet_fonts if f.startswith(_BAD_FONTS))
    add("No symbol fonts (Wingdings/Symbol) that parse as junk", not bad, ", ".join(bad) or "clean")
    pua = sorted({hex(ord(ch)) for ch in text if 0xE000 <= ord(ch) <= 0xF8FF})
    add("No private-use characters", not pua, ", ".join(pua[:5]) or "clean")
    date_re = re.compile(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4}\b")
    add("Dates in 'Mon YYYY' format", len(date_re.findall(text)) >= 2, f"{len(date_re.findall(text))} found")
    words = len(re.findall(r"\b\w+\b", text))
    add("Length fits ~2 pages (500-950 words)", 500 <= words <= 950, f"{words} words", warn=True)
    header_text = " ".join(p.text for s in d.sections for p in s.header.paragraphs).strip()
    add("Nothing important hidden in page header", not header_text, header_text[:60] or "empty", warn=True)
    fails = [c for c in checks if c["level"] == "fail"]
    return {"pass": not fails, "checks": checks, "words": words}
