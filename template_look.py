"""Learn the look of the candidate's tagged resume so drafts copy it.

Reads the template .docx paragraph by paragraph and finds: the name line, the contact line, section
headings, the headline, summary paragraphs, bullets (highlights, job bullets, a labeled "flagship"
bullet), and the job lines (employer line, job title line, scope line). For each it keeps a copy of
the paragraph settings (spacing, borders, shading, tabs, indents, style, bullet numbering) and the
text settings (font, size, color, bold, italics, caps). The builder then writes new words with those
exact settings, under the template's own section names and in its order.

If the template uses tables, text boxes or more than one column (these also confuse hiring software),
or has no recognizable Experience heading, learn() returns None and a plain-language reason, and the
builder uses its built-in design instead.
"""
import copy
import re

from docx.oxml.ns import qn

BAD_FONTS = ("wingdings", "symbol", "webdings", "zapf")
SAFE_BULLET = "•"
_BULLET_CHARS = "•▪■●◦‣∙·-–—*>Ø§§o"   # typed bullets; Ø and § are what Wingdings arrows/squares look like as text

# section kinds, checked in this order (first match wins)
SECTION_WORDS = [
    ("additional", ("additional",)),
    ("tech", ("technolog", "technical", "tools", "software")),
    ("recognitions", ("recogni", "award", "honor", "honour")),
    ("education", ("education", "academic", "degree")),
    ("certifications", ("certif", "license", "licence", "training")),
    ("experience", ("experience", "employment", "career history", "work history", "professional background",
                    "leadership history")),
    ("competencies", ("competenc", "skill", "expertise", "strengths", "capabilit", "areas of")),
    ("summary", ("summary", "profile", "overview", "about", "objective", "qualifications")),
]

_DATE_RE = re.compile(r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4}\b|\b(?:19|20)\d{2}\b"
                      r"|\bPresent\b", re.I)
_CONTACT_RE = re.compile(r"@|linkedin|\(?\d{3}\)?[\s.\-]\d{3}[\s.\-]\d{4}", re.I)


def section_kind(text):
    t = " ".join(text.lower().replace("&", " and ").split()).strip(" :")
    if not t or len(t) > 50 or len(t.split()) > 6:
        return None
    if t.startswith("additional"):
        return "additional"
    for kind, words in SECTION_WORDS:
        if any(w in t for w in words):
            return kind
    return None


def _bad_font(name):
    return (name or "").lower().startswith(BAD_FONTS)


def _ppr(p):
    return p._p.find(qn("w:pPr"))


def _text_runs(p):
    return [r for r in p.runs if r.text]


def _rpr_copy(run):
    """Copy of a run's text settings, minus symbol fonts and hidden text."""
    if run is None:
        return None
    rpr = run._r.find(qn("w:rPr"))
    if rpr is None:
        return None
    rpr = copy.deepcopy(rpr)
    for el in rpr.findall(qn("w:vanish")):
        rpr.remove(el)
    rf = rpr.find(qn("w:rFonts"))
    if rf is not None and any(_bad_font(rf.get(qn(a))) for a in ("w:ascii", "w:hAnsi", "w:cs")):
        rpr.remove(rf)
    return rpr


def _ppr_copy(p, keep_numbering=True):
    ppr = _ppr(p)
    if ppr is None:
        return None
    ppr = copy.deepcopy(ppr)
    for tag in ("w:sectPr", "w:pPrChange"):  # never copy a section break or tracked change
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    if not keep_numbering:
        for el in ppr.findall(qn("w:numPr")):
            ppr.remove(el)
    return ppr


def _has(p, path):
    ppr = _ppr(p)
    return ppr is not None and ppr.find(qn(path)) is not None


def _style_numbered(p):
    st = p.style
    while st is not None:
        sppr = st.element.find(qn("w:pPr"))
        if sppr is not None and sppr.find(qn("w:numPr")) is not None:
            return True
        st = st.base_style
    return False


def _bullet_prefix(p):
    """Leading typed bullet ('• ', '-\t', a Wingdings arrow...) -> (prefix_text, run) or (None, None)."""
    runs = _text_runs(p)
    if not runs:
        return None, None
    t = runs[0].text
    m = re.match(r"^\s*([" + re.escape(_BULLET_CHARS) + r"-])([\s\t]+)", t)
    if not m and len(t.strip()) == 1 and (t.strip() in _BULLET_CHARS or 0xF000 <= ord(t.strip()) <= 0xF0FF):
        m = re.match(r"^\s*(.)(\s*)$", t)
    if not m:
        return None, None
    ch, gap = m.group(1), m.group(2) or " "
    rf = runs[0]._r.find(qn("w:rPr") + "/" + qn("w:rFonts"))
    sym = rf is not None and any(_bad_font(rf.get(qn(a))) for a in ("w:ascii", "w:hAnsi"))
    if sym or 0xF000 <= ord(ch) <= 0xF0FF or ch in "Ø§o":
        ch = SAFE_BULLET
    return ch + gap, runs[0]


def is_bullet(p):
    return _has(p, "w:numPr") or _style_numbered(p) or _bullet_prefix(p)[0] is not None


def _signature(p):
    runs = _text_runs(p)
    r = runs[0] if runs else None
    rp = r._r.find(qn("w:rPr")) if r is not None else None
    ppr = _ppr(p)

    def x(el, tag):
        f = el.find(qn(tag)) if el is not None else None
        return None if f is None else (f.get(qn("w:val")) or "on")
    return (p.style.name if p.style is not None else "", x(rp, "w:b"), x(rp, "w:caps"), x(rp, "w:sz"),
            x(rp, "w:color"), ppr is not None and ppr.find(qn("w:pBdr")) is not None,
            ppr is not None and ppr.find(qn("w:shd")) is not None)


def _looks_heading(p, text):
    if not text or len(text) > 60 or len(text.split()) > 7 or text.endswith((".", ",", ";")) or is_bullet(p):
        return False
    if _DATE_RE.search(text) or _CONTACT_RE.search(text):
        return False
    style = (p.style.name if p.style is not None else "").lower()
    return section_kind(text) is not None or style.startswith("heading")


class Look:
    """What was learned. `parts` maps a part name to {"ppr", "rpr", "rpr2", "prefix", "label"}."""

    def __init__(self):
        self.parts = {}
        self.sections = []          # [(kind, heading text)] in template order
        self.name_with_contact = False

    def has(self, part):
        return part in self.parts

    def found(self):
        return sorted(self.parts)


def _part(p, keep_numbering=True, second_run=False, main=False):
    runs = _text_runs(p)
    prefix, prun = _bullet_prefix(p)
    body_runs = [r for r in runs if r is not prun] if prefix else runs
    if prefix and prun is not None:
        rest = prun.text.lstrip()[len(prefix.strip()):].lstrip()
        if rest:
            body_runs = [prun] + body_runs
    first = body_runs[0] if body_runs else (runs[0] if runs else None)
    if main and body_runs:   # paragraphs of running text: copy the look of most of the words
        first = max(body_runs, key=lambda r: len(r.text))
    second = None
    if second_run and len(body_runs) > 1:
        # the run after a tab (dates on the right), else the last run
        after_tab = [r for i, r in enumerate(body_runs) if i and ("\t" in body_runs[i - 1].text or r.text.startswith("\t"))]
        second = after_tab[0] if after_tab else body_runs[-1]
    out = {"ppr": _ppr_copy(p, keep_numbering), "rpr": _rpr_copy(first), "rpr2": _rpr_copy(second),
           "prefix": prefix, "prefix_rpr": _rpr_copy(prun) if prefix else None, "label": None, "src": p._p}
    return out


def _label(p, any_look=False):
    """'Key Achievement: did X' -> (label text, label rPr, text rPr).

    Counts when the label looks different from the text, or (any_look) when the bullet is shaded."""
    runs = _text_runs(p)
    prefix, prun = _bullet_prefix(p)
    runs = [r for r in runs if r is not prun]
    if not runs:
        return None
    text = "".join(r.text for r in runs).strip()
    m = re.match(r"^([A-Z][A-Za-z/&' -]{2,40}):\s+\S", text)
    if not m or len(m.group(1).split()) > 4:
        return None
    label = m.group(1) + ": "
    after = [r for r in runs[1:] if r.text.strip()]
    first_is_label = runs[0].text.strip().endswith(":") and len(runs) > 1
    differs = first_is_label and after and _signature_run(runs[0]) != _signature_run(after[0])
    if not (differs or any_look):
        return None
    return label, _rpr_copy(runs[0]), _rpr_copy(after[0] if first_is_label and after else runs[0])


def _signature_run(r):
    rp = r._r.find(qn("w:rPr"))
    return "" if rp is None else rp.xml


def unsupported(doc):
    """Plain-language reason the template can't be copied, or ''."""
    body = doc.element.body
    if body.findall(qn("w:tbl")):
        return "it uses a table"
    if body.xpath(".//w:txbxContent"):
        return "it uses text boxes"
    for cols in body.xpath(".//w:sectPr/w:cols"):
        if int(cols.get(qn("w:num")) or "1") > 1:
            return "it uses more than one column"
    return ""


def learn(doc):
    """Return (Look, "") or (None, reason)."""
    why = unsupported(doc)
    if why:
        return None, why
    paras = [p for p in doc.paragraphs]
    look = Look()
    texts = [" ".join(p.text.split()) for p in paras]

    # headings: section words, then anything that looks exactly like a found heading
    heads = {i for i, (p, t) in enumerate(zip(paras, texts)) if _looks_heading(p, t)}
    if heads:
        sigs = {_signature(paras[i]) for i in heads}
        for i, (p, t) in enumerate(zip(paras, texts)):
            if (i not in heads and t and len(t) <= 50 and not is_bullet(p) and not _DATE_RE.search(t)
                    and not _CONTACT_RE.search(t) and _signature(p) in sigs and t.upper() == t):
                heads.add(i)
    first_text = next((i for i, t in enumerate(texts) if t), None)
    heads.discard(first_text)   # the first line is the name, even if styled like a heading
    if not any(section_kind(texts[i]) == "experience" for i in heads):
        return None, "no Experience heading was found in it"
    for i in sorted(heads):
        kind = section_kind(texts[i]) or "other"
        look.sections.append((kind, texts[i].rstrip(":").strip() if kind != "other" else texts[i]))
    look.parts["heading"] = _part(paras[min(i for i in heads if section_kind(texts[i]))], keep_numbering=False)

    # walk the document section by section
    section = "top"
    exp_lines = []           # non-bullet lines in the first job block
    exp_seen_bullet = False
    first_of = {}
    for i, (p, t) in enumerate(zip(paras, texts)):
        if not t:
            continue
        if i in heads:
            section = section_kind(t) or "other"
            if section == "additional" and "additional_label" not in look.parts:
                look.parts["additional_label"] = dict(_part(p, keep_numbering=False), text=t)
                section = "experience"
            continue
        bullet = is_bullet(p)
        if i == first_text:
            look.parts["name"] = _part(p, keep_numbering=False, second_run=True)
            if _CONTACT_RE.search(t):
                look.name_with_contact = True
            continue
        if section in ("top", "summary"):
            if _CONTACT_RE.search(t) and not bullet and "contact" not in look.parts and len(t) < 160:
                look.parts["contact"] = _part(p, keep_numbering=False)
            elif bullet:
                look.parts.setdefault("highlight", _part(p))
            elif len(t) < 120 and not t.endswith(".") and "summary" not in look.parts:
                look.parts.setdefault("headline", _part(p, keep_numbering=False))
            else:
                look.parts.setdefault("summary", _part(p, keep_numbering=False, main=True))
            continue
        if section == "experience":
            if bullet:
                exp_seen_bullet = True
                ppr = _ppr(p)
                shaded = ppr is not None and ppr.find(qn("w:shd")) is not None
                lab = _label(p, any_look=shaded)
                if (lab or shaded) and "flagship" not in look.parts:
                    look.parts["flagship"] = _part(p)
                    if lab:
                        look.parts["flagship"]["label"] = lab[:2]
                        look.parts["flagship"]["rpr"] = lab[2]
                else:
                    look.parts.setdefault("bullet", _part(p))
            elif t.endswith(":") and len(t.split()) <= 5:
                look.parts.setdefault("additional_label", dict(_part(p, keep_numbering=False), text=t))
            elif not exp_seen_bullet:
                exp_lines.append(p)
            elif "additional_role" not in look.parts and exp_seen_bullet and "additional_label" in look.parts:
                look.parts["additional_role"] = _part(p, keep_numbering=False, second_run=True)
            continue
        if section not in first_of:
            first_of[section] = p
    # the first job block: short lines are employer / title, a long one is the scope line
    short = [p for p in exp_lines if len(p.text.strip()) < 110 and not p.text.strip().endswith(".")]
    long_ = [p for p in exp_lines if p not in short]
    if short:
        look.parts["org"] = _part(short[0], keep_numbering=False, second_run=True)
        look.parts["title"] = _part(short[1] if len(short) > 1 else short[0], keep_numbering=False, second_run=True)
    if long_:
        look.parts["scope"] = _part(long_[0], keep_numbering=False, main=True)
    for kind, p in first_of.items():
        look.parts["sec_" + kind] = _part(p, main=True)
    if "bullet" not in look.parts:
        if "highlight" in look.parts:
            look.parts["bullet"] = look.parts["highlight"]
        elif "flagship" in look.parts:
            look.parts["bullet"] = dict(look.parts["flagship"], label=None)
    if "bullet" not in look.parts:
        return None, "no bullet points were found in its Experience section"
    look.parts.setdefault("highlight", look.parts["bullet"])
    body = look.parts.get("summary") or look.parts.get("scope") or {"ppr": None, "rpr": None, "rpr2": None,
                                                                     "prefix": None, "prefix_rpr": None, "label": None}
    look.parts.setdefault("body", body)
    return look, ""


def sanitize_bullets(doc):
    """Make every bullet definition ATS-safe: plain '•' in a normal font, never Wingdings or picture bullets."""
    try:
        numbering = doc.part.numbering_part.element
    except (KeyError, NotImplementedError):
        return 0
    fixed = 0
    for lvl in numbering.xpath("./w:abstractNum/w:lvl"):
        fmt = lvl.find(qn("w:numFmt"))
        if fmt is None or fmt.get(qn("w:val")) != "bullet":
            continue
        txt = lvl.find(qn("w:lvlText"))
        rf = lvl.find(qn("w:rPr") + "/" + qn("w:rFonts"))
        pic = lvl.find(qn("w:lvlPicBulletId"))
        val = txt.get(qn("w:val")) if txt is not None else ""
        bad = (pic is not None or not val or any(0xE000 <= ord(c) <= 0xF8FF for c in val)
               or (rf is not None and any(_bad_font(rf.get(qn(a))) for a in ("w:ascii", "w:hAnsi", "w:cs"))))
        if not bad:
            continue
        if pic is not None:
            lvl.remove(pic)
        if txt is not None:
            txt.set(qn("w:val"), SAFE_BULLET)
        if rf is not None:
            rpr = rf.getparent()
            rpr.remove(rf)
        fixed += 1
    return fixed


# ---------------------------------------------------------------- formatting report (no resume text)

def _fmt(p):
    runs = _text_runs(p)
    r = runs[0] if runs else None
    bits = []
    ppr = _ppr(p)
    if ppr is not None:
        for tag, name in (("w:pBdr", "line"), ("w:shd", "shade"), ("w:numPr", "list"), ("w:tabs", "tabstop"),
                          ("w:jc", "align")):
            el = ppr.find(qn(tag))
            if el is not None:
                if tag == "w:pBdr":
                    name += ":" + ",".join(c.tag.split("}")[1] for c in el)
                if tag == "w:jc":
                    name += ":" + (el.get(qn("w:val")) or "")
                if tag == "w:shd":
                    name += ":" + (el.get(qn("w:fill")) or "")
                bits.append(name)
    if _style_numbered(p) and "list" not in bits:
        bits.append("list(style)")
    pre, _ = _bullet_prefix(p)
    if pre:
        bits.append("typed-bullet")
    if r is not None:
        f = r.font
        bits.append(f"font={f.name or '-'} {f.size.pt if f.size else '-'}pt")
        for a in ("bold", "italic", "all_caps", "small_caps", "underline"):
            if getattr(f, a):
                bits.append(a)
        if f.color is not None and f.color.type is not None and f.color.rgb is not None:
            bits.append(f"color={f.color.rgb}")
        if len(runs) > 1:
            bits.append(f"{len(runs)} runs")
        if any("\t" in x.text for x in runs):
            bits.append("has-tab")
    for el in p._p.iter(qn("w:br")):
        bits.append("line-break")
        break
    return " ".join(bits)


def report(path):
    """Paragraph-by-paragraph formatting and what each was read as. Shows section headings; hides other text."""
    from docx import Document
    doc = Document(path)
    look, why = learn(doc)
    out = [f"file: {path.rsplit('/', 1)[-1]}", f"result: {'COPY TEMPLATE' if look else 'BUILT-IN (' + why + ')'}"]
    used = {}
    if look:
        out.append("sections: " + " > ".join(k for k, _ in look.sections))
        out.append("found: " + ", ".join(look.found()))
        for name, part in look.parts.items():
            if part.get("src") is not None:
                used.setdefault(part["src"], set()).add(name)
    for i, p in enumerate(doc.paragraphs):
        t = " ".join(p.text.split())
        if not t:
            continue
        shown = t if (section_kind(t) and len(t) <= 50) else f"[{len(t.split())} words]"
        role = ",".join(sorted(used.get(p._p, []))) or "-"
        style = p.style.name if p.style is not None else "-"
        out.append(f"{i:3d} {role:<22} style={style} | {_fmt(p)} | {shown}")
    return "\n".join(out)


if __name__ == "__main__":
    import sys
    for a in sys.argv[1:]:
        print(report(a))
        print()
